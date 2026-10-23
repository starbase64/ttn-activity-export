"""Read-only retrieval, resumable account-isolated caching and explicit export formats."""
from __future__ import annotations
import csv
import fnmatch
import hashlib
import inspect
import json
import os
import re
import threading
from dataclasses import dataclass, field
from datetime import date, timedelta, datetime, timezone
from pathlib import Path
from typing import Any, Callable
from catalog import DATASETS
from failures import diagnostic, error_status
from activity_table import table_row, type_graph, COLUMNS, SOURCE_PATHS

SCHEMA = 1
class Cancelled(Exception): pass
class ActivityPaginationError(Exception): pass


def request_spec(key,method,args=(),kwargs=None,binary_format=''):
    request={'schema':SCHEMA,'method':method,'args':args,
             'kwargs':{k:(v.name if hasattr(v,'name') else v) for k,v in (kwargs or {}).items()},'binary':binary_format}
    if key=='activities':request['reader']='unfiltered-paged-v2'
    return request


def fingerprint(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:20]


def account_key(email: str, china: bool = False) -> str:
    return fingerprint([email.strip().casefold(), china])


def atomic_json(path: Path, value: Any):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    os.replace(tmp, path)


def days(start: str, end: str):
    a, b = date.fromisoformat(start), date.fromisoformat(end)
    if a > b:
        raise ValueError('Start > end')
    if b > date.today():
        raise ValueError('End date is in the future')
    while a <= b:
        yield a.isoformat()
        a += timedelta(days=1)


def pointer_part(key):
    return str(key).replace('~', '~0').replace('/', '~1')


def leaves(data: Any, path=''):
    """Lossless JSON pointer rows, including null and empty containers."""
    if isinstance(data, dict) and data:
        for key, value in data.items():
            yield from leaves(value, path + '/' + pointer_part(key))
    elif isinstance(data, list) and data:
        for idx, value in enumerate(data):
            yield from leaves(value, path + '/' + str(idx))
    else:
        kind = ('null' if data is None else 'boolean' if isinstance(data, bool) else
                'number' if isinstance(data, (int, float)) else 'object' if isinstance(data, dict) else
                'array' if isinstance(data, list) else 'string')
        yield path or '/', data, kind


def field_patterns(data: Any, path=''):
    """Generalize record indices but retain scalar tuple positions (timestamp/value)."""
    if isinstance(data, dict) and data:
        for key, value in data.items():
            yield from field_patterns(value, path + '/' + pointer_part(key))
    elif isinstance(data, list) and data:
        nested = any(isinstance(value, (dict, list)) for value in data)
        for idx, value in enumerate(data):
            yield from field_patterns(value, path + '/' + ('*' if nested else str(idx)))
    else:
        yield path or '/'


def field_selected(path, selectors):
    if not selectors:
        return True
    # Match tokens, not arbitrary slash-crossing globs. A selected parent includes descendants.
    tokens = path.strip('/').split('/')
    for selector in selectors:
        parts = selector.strip('/').split('/')
        if len(parts) <= len(tokens) and all(fnmatch.fnmatchcase(t, p) for t, p in zip(tokens, parts)):
            return True
    return False


def safe_cell(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list, bool)):
        value = json.dumps(value, ensure_ascii=False)
    # Preserve numbers as numbers; protect spreadsheet consumers from formula injection.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def flat_record(record, prefix=''):
    if not isinstance(record, dict):
        return {'value': record}
    result = {}
    for key, val in record.items():
        path = prefix + '/' + pointer_part(key)
        if isinstance(val, dict) and val:
            result.update(flat_record(val, path))
        else:
            result[path] = val
    return result


@dataclass
class Plan:
    start: str
    end: str
    datasets: list[str]
    output: str
    account: str
    formats: list[str] = field(default_factory=lambda: ['csv'])
    layout: str = 'long'
    delimiter: str = ';'
    split_year: bool = True
    refresh_days: int = 2
    delay: float = 1.0
    retries: int = 2
    activity_type: str = ''
    activity_details: list[str] = field(default_factory=list)
    activity_files: list[str] = field(default_factory=list)
    selectors: list[str] = field(default_factory=list)
    expert_method: str = ''
    expert_kwargs: dict = field(default_factory=dict)
    expert_daily: bool = True
    activity_summary: bool = True
    language: str = 'de'
    preview_rows: int = 10
    preview_requests: int = 12
    source: str = 'online'
    archive_path: str = ''
    optimize: bool = True
    activity_batch_days: int = 31
    health_detail: str = 'full'
    health_batch_days: int = 28

    def validate(self):
        date.fromisoformat(self.start); date.fromisoformat(self.end)
        if self.start > self.end or self.end > date.today().isoformat():
            raise ValueError('Invalid date range')
        if not self.datasets and not self.expert_method:
            raise ValueError('Select at least one dataset')
        if not self.formats or not set(self.formats) <= {'csv', 'json', 'xlsx'}:
            raise ValueError('Select an output format')
        if not 1<=self.preview_rows<=50 or not 1<=self.preview_requests<=50:
            raise ValueError('Invalid preview limits')
        if not isinstance(self.optimize,bool) or not 1<=self.activity_batch_days<=365:
            raise ValueError('Invalid transfer options')
        if self.health_detail not in ('full','daily') or type(self.health_batch_days) is not int or not 1<=self.health_batch_days<=28:
            raise ValueError('Invalid health transfer options')
        if self.source not in ('online','offline') or (self.source=='offline' and not self.archive_path):
            raise ValueError('Invalid offline source')
        if self.language not in ('de','en'):
            raise ValueError('Invalid output language')
        if self.layout not in ('long', 'wide') or self.delimiter not in (';', ',', '\t'):
            raise ValueError('Invalid table settings')
        if not 0 <= self.delay <= 60 or not 0 <= self.retries <= 5 or self.refresh_days < 0:
            raise ValueError('Invalid request settings')
        if not re.fullmatch(r'[a-f0-9]{20}', self.account):
            raise ValueError('Invalid account key')
        if not set(self.datasets) <= {d.key for d in DATASETS}:
            raise ValueError('Unknown dataset')
        from catalog import ACTIVITY_DETAILS
        if not set(self.activity_details) <= set(ACTIVITY_DETAILS):
            raise ValueError('Unknown activity detail')
        if not set(self.activity_files) <= {'ORIGINAL', 'GPX', 'TCX', 'KML', 'CSV'}:
            raise ValueError('Unknown activity format')
        if self.expert_method and not is_read_method(self.expert_method):
            raise ValueError('Only public get_* methods are supported')
        if not isinstance(self.expert_kwargs, dict):
            raise ValueError('Parameters must be a JSON object')
        if self.expert_kwargs.get('generate_defaults'):
            raise ValueError('generate_defaults is not allowed')
        if (self.activity_details or self.activity_files) and 'activities' not in self.datasets:
            raise ValueError('Select activities to export activity details/files')


def is_read_method(name):
    return bool(re.fullmatch(r'get_[a-zA-Z0-9_]+', name))


def read_methods(api_class):
    return {name: str(inspect.signature(fn)) for name, fn in inspect.getmembers(api_class, inspect.isfunction)
            if is_read_method(name)}


def substitute(value, day, start, end):
    if isinstance(value, str):
        return value.replace('{date}', day).replace('{start}', start).replace('{end}', end)
    if isinstance(value, list):
        return [substitute(v, day, start, end) for v in value]
    if isinstance(value, dict):
        return {k: substitute(v, day, start, end) for k, v in value.items()}
    return value


class Exporter:
    def __init__(self, api, plan: Plan, cancel=None, event: Callable = lambda *a: None):
        plan.validate()
        self.api, self.plan = api, plan
        self.cancel = cancel or threading.Event()
        self.event = event
        self.root = Path(plan.output).expanduser().resolve() / ('account-' + plan.account)
        self.cache = self.root / 'cache-v1'
        self.stats = {'downloaded': 0, 'cached': 0, 'errors': 0, 'empty': 0, 'activities': 0}
        self.records = []
        self.errors = []
        self.last_request = 0.0
        self.type_descendants = None
        self.warnings = []
        if plan.health_detail=='daily':self.warnings.append('daily_statistics_without_timelines')
        self._staged={};self.eta=None;self.transfer={'requests':0,'range_calls':0,'fallback_days':0}
        self.range_disabled=set()

    def cache_path(self,key,day,method,args=(),kwargs=None,binary_format=''):
        suffix={'ORIGINAL':'.zip','GPX':'.gpx','TCX':'.tcx','KML':'.kml','CSV':'.csv'}.get(binary_format,'.json')
        return self.cache/key/day[:4]/(day+'-'+fingerprint(request_spec(key,method,args,kwargs,binary_format))+suffix)

    def refresh_day(self,day):
        return day=='snapshot' or (self.plan.refresh_days>0 and date.fromisoformat(day)>=date.today()-timedelta(days=self.plan.refresh_days-1))

    def available(self,key,day,method,args=(),kwargs=None,binary_format=''):
        if self.refresh_day(day):return False
        path=self.cache_path(key,day,method,args,kwargs,binary_format)
        if not path.is_file():return False
        try:
            if binary_format:return path.stat().st_size>0
            stored=json.loads(path.read_text(encoding='utf-8'))
            return 'data' in stored and stored['request']==json.loads(json.dumps(request_spec(key,method,args,kwargs,binary_format)))
        except (ValueError,KeyError,OSError):return False

    def check(self):
        if self.cancel.is_set():
            raise Cancelled()

    def wait(self, seconds):
        if self.cancel.wait(seconds):
            raise Cancelled()

    def call(self, fn, *args, **kwargs):
        import time
        from garminconnect import (GarminConnectAuthenticationError, GarminConnectTooManyRequestsError,
                                   GarminConnectConnectionError, GarminConnectNotFoundError)
        for attempt in range(self.plan.retries + 1):
            self.check()
            request_began=time.monotonic()
            self.wait(max(0.0, self.plan.delay - (time.monotonic() - self.last_request)))
            self.last_request = time.monotonic()
            self.transfer['requests']+=1;self.event('request',self.transfer['requests'])
            try:
                return fn(*args, **kwargs)
            except (GarminConnectAuthenticationError, GarminConnectTooManyRequestsError):
                raise  # Stop immediately; never hammer a rate-limited account.
            except GarminConnectNotFoundError:
                raise
            except GarminConnectConnectionError as exc:
                if error_status(exc) in (400,401,403,404,405,422,429):
                    raise
                if attempt == self.plan.retries:
                    raise
                pause=min(30,2**(attempt+1));self.event('retry_wait',pause)
                self.wait(pause);self.event('retry_wait',0)
            finally:
                if self.eta:self.eta.observe_request(time.monotonic()-request_began)

    def retrieve(self, key, day, method, args=(), kwargs=None, binary_format=''):
        self.check()
        kwargs = kwargs or {}
        request = request_spec(key,method,args,kwargs,binary_format)
        digest = fingerprint(request)
        suffix = {'ORIGINAL': '.zip', 'GPX': '.gpx', 'TCX': '.tcx', 'KML': '.kml', 'CSV': '.csv'}.get(binary_format, '.json')
        path = self.cache / key / day[:4] / (day + '-' + digest + suffix)
        refresh = day != 'snapshot' and self.plan.refresh_days > 0 and date.fromisoformat(day) >= date.today() - timedelta(days=self.plan.refresh_days - 1)
        reuse = path.exists() and not refresh and day != 'snapshot'
        payload = None
        if reuse and not binary_format:
            try:
                stored = json.loads(path.read_text(encoding='utf-8'))
                if stored['request'] != json.loads(json.dumps(request)):
                    reuse = False
                else:
                    payload = stored['data']
            except (ValueError, KeyError, OSError):
                reuse = False
        if reuse:
            self.stats['cached'] += 1
        else:
            try:
                fn = getattr(self.api, method)
                if (key,day) in self._staged:payload=self._staged.pop((key,day))
                else:payload = self.fetch_data(key,day,method,args,kwargs)
                self.check()
                if binary_format:
                    if not isinstance(payload, bytes) or not payload:
                        raise ValueError('Empty or invalid download')
                    path.parent.mkdir(parents=True, exist_ok=True)
                    tmp = path.with_suffix(path.suffix + '.tmp')
                    tmp.write_bytes(payload)
                    os.replace(tmp, path)
                else:
                    atomic_json(path, {'request': request, 'fetched_utc': datetime.now(timezone.utc).isoformat(), 'data': payload})
                self.stats['downloaded'] += 1
            except Cancelled:
                raise
            except Exception as exc:
                from garminconnect import GarminConnectAuthenticationError, GarminConnectTooManyRequestsError
                if isinstance(exc, (GarminConnectAuthenticationError, GarminConnectTooManyRequestsError)):
                    raise
                # Never persist arbitrary exception text (it can contain sensitive response data).
                self.stats['errors'] += 1
                self.errors.append({'dataset': key, 'date': day, 'error_type': type(exc).__name__, 'http_status': error_status(exc)})
                self.event('error', key, day, type(exc).__name__ + (f' · HTTP {error_status(exc)}' if error_status(exc) else ''))
                from garminconnect import GarminConnectConnectionError, GarminConnectNotFoundError
                if isinstance(exc,GarminConnectConnectionError) and not isinstance(exc,GarminConnectNotFoundError):
                    raise
                return None
        if payload in (None, [], {}) and not binary_format:
            self.stats['empty'] += 1
        self.records.append({'dataset': key, 'date': day, 'path': str(path.relative_to(self.root)), 'binary': bool(binary_format)})
        self.event('item', key, day, 'cached' if reuse else 'downloaded')
        return payload

    def fetch_data(self,key,day,method,args=(),kwargs=None):
        if key=='activities':return self.fetch_activity_day(day)
        payload=self.call(getattr(self.api,method),*args,**(kwargs or {}))
        if key.endswith('-daily'):
            from transfer import split_dated_records
            # Only a successful, explicitly empty single-day response means empty.
            if payload==[]:return []
            grouped=split_dated_records(payload,[day])
            if grouped is None or day not in grouped:raise ValueError('Invalid daily statistics response')
            return grouped[day]
        return payload

    def run(self):
        p = self.plan
        stopped = False
        fatal = ''
        fatal_detail = {}
        try:
            self.collect()
        except Cancelled:
            stopped = True
        except Exception as exc:
            fatal = type(exc).__name__
            fatal_detail = diagnostic(exc)
        self.event('writing')
        # Write only data actually retrieved in this run, never unrelated cache entries.
        config = {k: v for k, v in vars(p).items() if k not in ('account', 'output')}
        destination = self.root / 'exports' / (p.start + '_' + p.end + '_' + fingerprint(config))
        destination.mkdir(parents=True, exist_ok=True)
        export_error = ''
        try:
            self.write_outputs(destination)
        except Exception as exc:
            export_error = type(exc).__name__
        report = {'schema': SCHEMA, 'completed': not stopped and not fatal and not export_error and not self.stats['errors'],
                  'cancelled': stopped, 'fatal_error_type': fatal, 'fatal_diagnostic': fatal_detail, 'export_error_type': export_error,
                  'stats': self.stats, 'errors': self.errors, 'warnings': self.warnings, 'config': config,
                  'records': self.records, 'transfer':self.transfer, 'created_utc': datetime.now(timezone.utc).isoformat()}
        atomic_json(destination / 'report.json', report)
        self.event('eta_end','cancelled' if stopped else 'error' if fatal or export_error or self.stats['errors'] else 'complete')
        return str(destination), report

    def collect(self):
        from transfer import plan_jobs,RemainingTime
        p=self.plan
        from catalog import selected_datasets
        selected=selected_datasets(p)
        total_days=(date.fromisoformat(p.end)-date.fromisoformat(p.start)).days+1
        self.event('total',sum(1 if d.kind=='account' else total_days for d in selected)
                   +(total_days if p.expert_daily else 1)*bool(p.expert_method))
        self.event('phase','planning');jobs=plan_jobs(self,selected)
        expert_count=(total_days if p.expert_daily else 1)*bool(p.expert_method)
        if p.source!='offline':
            self.eta=RemainingTime([job.family for job in jobs]+['expert']*expert_count,p.delay,self.event)
            self.eta.open_activity_days=sum(len(job.days) for job in jobs if job.dataset.key=='activities')
        self.event('phase','download' if p.source!='offline' else 'offline')
        for job in jobs:
            self.check();ds=job.dataset
            if self.eta:self.eta.begin(job.family)
            if job.mode=='range' and ds.key not in self.range_disabled:self.prefetch_range(ds,job.days)
            for day in job.days:
                args = () if ds.kind == 'account' else (day, day) if ds.kind == 'range' else (day,)
                kwargs = {}
                if ds.key == 'endurance':
                    kwargs['aggregation'] = 'daily'
                result = self.retrieve(ds.key, day, ds.method, args, kwargs)
                self.event('progress')
                if ds.key == 'activities' and isinstance(result, list):
                    matched=self.match_activities(result)
                    self.stats['activities']+=len(matched)
                    # Finish the list job before adding its separately measured extras.
                    if self.eta and self.eta.active:self.eta.finish()
                    self.activity_exports(day, matched)
                    if self.eta:self.eta.open_activity_days-=1;self.eta.emit()
            if self.eta and self.eta.active:self.eta.finish()
        if p.expert_method:
            if not callable(getattr(self.api, p.expert_method, None)):
                raise ValueError('Unsupported method')
            for day in (days(p.start, p.end) if p.expert_daily else [p.start]):
                if self.eta:self.eta.begin('expert')
                kwargs = substitute(p.expert_kwargs, day, p.start, p.end)
                inspect.signature(getattr(self.api, p.expert_method)).bind(**kwargs)
                self.retrieve('expert-' + p.expert_method, day, p.expert_method, kwargs=kwargs)
                self.event('progress')
                if self.eta:self.eta.finish()

    def prefetch_range(self,ds,calendar):
        self.event('batch',ds.key,calendar[0],calendar[-1]);self.transfer['range_calls']+=1
        try:
            if ds.key=='activities':
                data=self.fetch_activity_range(calendar[0],calendar[-1]);grouped={day:[] for day in calendar}
                for a in data:
                    stamp=a.get('startTimeLocal')
                    if not isinstance(stamp,str) or stamp[:10] not in grouped:
                        raise ActivityPaginationError('Activity range date mismatch')
                    grouped[stamp[:10]].append(a)
            elif ds.key=='body_battery':
                from transfer import split_body_battery
                grouped=split_body_battery(self.call(getattr(self.api,ds.method),calendar[0],calendar[-1]),calendar)
                if grouped is None:raise ValueError('Unexpected range shape')
            else:
                from transfer import split_dated_records
                summary=ds.key=='summary'
                method='get_daily_summaries' if summary else ds.method
                payload=self.call(getattr(self.api,method),calendar[0],calendar[-1])
                grouped=split_dated_records(payload,calendar,summary=summary)
                if grouped is None:raise ValueError('Unexpected dated range shape')
            for day,payload in grouped.items():self._staged[(ds.key,day)]=payload
            self.transfer['fallback_days']+=len(calendar)-len(grouped)
        except Cancelled:raise
        except Exception as exc:
            from garminconnect import GarminConnectConnectionError,GarminConnectAuthenticationError,GarminConnectTooManyRequestsError
            status=error_status(exc)
            if isinstance(exc,(GarminConnectAuthenticationError,GarminConnectTooManyRequestsError)):raise
            if isinstance(exc,GarminConnectConnectionError) and status not in (400,404,405,422):raise
            self.range_disabled.add(ds.key);self.transfer['fallback_days']+=len(calendar)
            self.warnings.append('range_fallback:'+ds.key+':'+type(exc).__name__)
            self.event('batch_fallback',ds.key)

    def fetch_activity_range(self,start,end):
        items=[];seen=set();offset=0
        for _ in range(2000):
            batch=self.call(self.api.search_activities,start=offset,limit=100,startDate=start,endDate=end)
            if not isinstance(batch,list):raise ActivityPaginationError('Unexpected activity response')
            if not batch:return items
            added=0
            for a in batch:
                if not isinstance(a,dict):raise ActivityPaginationError('Invalid activity')
                token=str(a['activityId']) if a.get('activityId') is not None else fingerprint(a)
                if token not in seen:items.append(a);seen.add(token);added+=1
            if not added:raise ActivityPaginationError('Repeated activity page')
            offset+=len(batch)
        raise ActivityPaginationError('Activity pagination limit exceeded')

    def fetch_activity_day(self,day):
        # Never send a subtype such as e_bike_mountain as the server's activityType.
        # Explicit pagination allows rate limiting and cancellation between pages.
        search=getattr(self.api,'search_activities',None)
        if not callable(search):
            return self.call(self.api.get_activities_by_date,day,day)
        items=[];seen=set();offset=0
        for _ in range(2000):
            batch=self.call(search,start=offset,limit=100,startDate=day,endDate=day)
            if not isinstance(batch,list):
                raise ActivityPaginationError('Unexpected activity response')
            if not batch:return items
            added=0
            for activity in batch:
                if not isinstance(activity,dict):
                    raise ActivityPaginationError('Invalid activity record')
                ident=activity.get('activityId')
                token=str(ident) if ident is not None else fingerprint(activity)
                if token not in seen:
                    seen.add(token);items.append(activity);added+=1
            if not added:
                raise ActivityPaginationError('Server repeated an activity page')
            offset+=len(batch)
        raise ActivityPaginationError('Activity pagination limit exceeded')

    def match_activities(self,activities):
        selected=self.plan.activity_type
        if not selected:return activities
        if self.type_descendants is None:
            self.type_descendants={}
            getter=getattr(self.api,'get_activity_types',None)
            if callable(getter):
                from garminconnect import GarminConnectAuthenticationError,GarminConnectTooManyRequestsError
                try:
                    self.type_descendants=type_graph(self.call(getter))
                except (GarminConnectAuthenticationError,GarminConnectTooManyRequestsError):
                    raise
                except Cancelled:
                    raise
                except Exception:
                    self.warnings.append('activity_taxonomy_unavailable_exact_match_only')
                    self.event('warning','activity_taxonomy_unavailable_exact_match_only')
        keys=self.type_descendants.get(selected,set())|{selected}
        return [a for a in activities if isinstance(a.get('activityType'),dict) and a['activityType'].get('typeKey') in keys]

    def activity_exports(self, day, activities):
        from garminconnect import Garmin
        requests=[]
        for activity in activities:
            ident = activity.get('activityId')
            if not isinstance(ident, (int, str)) or not str(ident).isdigit():
                continue
            for method in self.plan.activity_details:
                requests.append(('activity-'+str(ident)+'-'+method.removeprefix('get_activity_'),method,(ident,),{},'', 'details'))
            for fmt in self.plan.activity_files:
                requests.append(('activity-'+str(ident)+'-'+fmt.lower(),'download_activity',(ident,),{'dl_fmt':Garmin.ActivityDownloadFormat[fmt]},fmt,'files'))
        self.event('total_extra',len(requests))
        classified=[]
        for key,method,args,kwargs,fmt,family in requests:
            family='cache' if self.available(key,day,method,args,kwargs,fmt) else family
            classified.append((key,method,args,kwargs,fmt,family))
            if self.eta:self.eta.add(family,1)
        for key,method,args,kwargs,fmt,family in classified:
            if self.eta:self.eta.begin(family)
            self.retrieve(key,day,method,args,kwargs,fmt);self.event('progress')
            if self.eta:self.eta.finish()

    def iter_payloads(self, records):
        for record in records:
            stored = json.loads((self.root / record['path']).read_text(encoding='utf-8'))
            payload=stored['data']
            if record['dataset']=='activities' and isinstance(payload,list):
                payload=self.match_activities(payload)
            yield record, payload

    def summary_fields(self):
        selectors=self.plan.selectors
        fields=[(de if self.plan.language=='de' else en) for key,de,en in COLUMNS
                if field_selected('/'+SOURCE_PATHS[key],selectors) or field_selected('/0/'+SOURCE_PATHS[key],selectors)]
        if not fields:
            raise ValueError('No selected fields exist in the readable activity summary')
        return fields

    def rows(self, records):
        for record, payload in self.iter_payloads(records):
            if record['dataset']=='activities' and self.plan.activity_summary:
                for activity in payload or []:
                    fields=self.summary_fields()
                    yield {k:v for k,v in table_row(activity,record['date'],self.plan.language).items() if k in fields}
            elif self.plan.layout == 'long':
                for pointer, value, kind in leaves(payload):
                    if field_selected(pointer, self.plan.selectors):
                        yield {'source_date': record['date'], 'dataset': record['dataset'],
                               'path': pointer, 'value': value, 'type': kind}
            else:
                for idx, item in enumerate(payload if isinstance(payload, list) else [payload]):
                    row = {'source_date': record['date'], 'dataset': record['dataset'], 'record_index': idx}
                    row.update({k: v for k, v in flat_record(item).items() if field_selected(k, self.plan.selectors)})
                    yield row

    def table_cells(self,row,is_activity_summary=False):
        cells={k:safe_cell(v) for k,v in row.items()}
        if is_activity_summary and self.plan.language=='de' and self.plan.delimiter in (';', '\t'):
            cells={k:(str(v).replace('.',',') if isinstance(v,float) else v) for k,v in cells.items()}
        return cells

    def copy_binary(self,record,target):
        import shutil
        shutil.copyfile(self.root/record['path'],target)

    def write_outputs(self, destination):
        import shutil
        groups = {}
        for record in self.records:
            year = record['date'][:4] if self.plan.split_year else 'all'
            key = record['dataset'] + '-' + year
            if record['binary']:
                target = destination / 'activities' / Path(record['path']).name
                # Source basename includes request hash and therefore activity identity.
                target.parent.mkdir(parents=True, exist_ok=True)
                self.copy_binary(record,target)
            else:
                groups.setdefault(key, []).append(record)
        for key, records in groups.items():
            if 'json' in self.plan.formats:
                # Stream daily envelopes, preserving native JSON types and nesting.
                target = destination / (key + '.json')
                tmp = target.with_suffix('.json.tmp')
                with tmp.open('w', encoding='utf-8') as f:
                    f.write('[\n')
                    for i, (record, payload) in enumerate(self.iter_payloads(records)):
                        if i: f.write(',\n')
                        json.dump({'source_date': record['date'], 'dataset': record['dataset'], 'data': payload}, f, ensure_ascii=False)
                    f.write('\n]\n')
                os.replace(tmp, target)
            if not set(self.plan.formats) & {'csv', 'xlsx'}:
                continue
            is_activity_summary=records[0]['dataset']=='activities' and self.plan.activity_summary
            fields = self.summary_fields() if is_activity_summary else (['source_date', 'dataset', 'path', 'value', 'type'] if self.plan.layout == 'long' else ['source_date', 'dataset', 'record_index'])
            if self.plan.layout == 'wide' and not is_activity_summary:
                found = set(fields)
                for row in self.rows(records):
                    for name in row:
                        if name not in found:
                            fields.append(name); found.add(name)
            if 'csv' in self.plan.formats:
                target = destination / (key + '.csv')
                tmp = target.with_suffix('.csv.tmp')
                with tmp.open('w', encoding='utf-8-sig', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=fields, delimiter=self.plan.delimiter)
                    writer.writeheader()
                    for row in self.rows(records):
                        writer.writerow(self.table_cells(row,is_activity_summary))
                os.replace(tmp, target)
            if 'xlsx' in self.plan.formats:
                from openpyxl import Workbook
                if len(fields) > 16384:
                    raise ValueError('Excel column limit; use CSV/JSON')
                book = Workbook(write_only=True)
                sheet = book.create_sheet('Data 1'); sheet.append(fields)
                n, page = 1, 1
                for row in self.rows(records):
                    if n >= 1048576:
                        page += 1; sheet = book.create_sheet('Data ' + str(page)); sheet.append(fields); n = 1
                    values = [safe_cell(row.get(k)) for k in fields]
                    if any(isinstance(v, str) and (len(v) > 32767 or re.search(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]', v)) for v in values):
                        raise ValueError('Value cannot be represented in Excel; use CSV/JSON')
                    sheet.append(values); n += 1
                target = destination / (key + '.xlsx')
                tmp = target.with_suffix('.tmp.xlsx'); book.save(tmp); os.replace(tmp, target)
