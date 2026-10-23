"""Local data browsing, activity-file import and portable archive backups.

No network access in this module. Imported records retain their original files.
"""
from __future__ import annotations
import csv
import hashlib
import io
import json
import math
import os
import sqlite3
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path
from archive import Archive, OfflineArchiveError
from engine import request_spec, safe_cell, Cancelled

MAX_FILE = 128 * 1024 * 1024
MAX_ZIP_TOTAL = 1024 * 1024 * 1024
MAX_ZIP_FILES = 30000

def valid_point(lat, lon):
    return isinstance(lat,(int,float)) and isinstance(lon,(int,float)) and math.isfinite(lat) and math.isfinite(lon) and -90 <= lat <= 90 and -180 <= lon <= 180

def local_name(tag): return tag.rsplit('}',1)[-1]

def xml_root(data):
    # A byte search misses UTF-16/32 encoded documents. Let expat decode the
    # document and refuse any DTD before ElementTree could expand entities.
    from xml.parsers import expat
    def forbidden(*args):raise ValueError('XML declarations are not supported')
    guard=expat.ParserCreate()
    guard.StartDoctypeDeclHandler=forbidden;guard.EntityDeclHandler=forbidden
    guard.ExternalEntityRefHandler=forbidden
    try:guard.Parse(data,True)
    except expat.ExpatError as exc:raise ET.ParseError(str(exc)) from exc
    return ET.fromstring(data)

def distance_m(points):
    total=0
    for a,b in zip(points,points[1:]):
        if a is None or b is None: continue
        lat1,lat2=map(math.radians,(a[0],b[0])); dl=math.radians(b[0]-a[0]);dn=math.radians(b[1]-a[1])
        h=math.sin(dl/2)**2+math.cos(lat1)*math.cos(lat2)*math.sin(dn/2)**2
        total+=6371000*2*math.asin(min(1,math.sqrt(h)))
    return total

def parse_xml_activity(data, fmt):
    root=xml_root(data); points=[];times=[]; name='';sport='other';duration=None;distance=None
    if fmt=='gpx':
        if local_name(root.tag)!='gpx':raise ValueError('Not GPX')
        for el in root.iter():
            key=local_name(el.tag)
            if key in ('trkseg','rte') and points and points[-1] is not None:points.append(None)
            if key=='name' and not name:name=el.text or ''
            if key=='type' and el.text:sport=el.text.strip().lower()
            if key in ('trkpt','rtept'):
                lat,lon=float(el.attrib['lat']),float(el.attrib['lon'])
                if valid_point(lat,lon):points.append((lat,lon))
                for child in el:
                    if local_name(child.tag)=='time' and child.text:times.append(child.text)
    else:
        if local_name(root.tag)!='TrainingCenterDatabase':raise ValueError('Not TCX')
        for el in root.iter():
            key=local_name(el.tag)
            if key=='Activity':sport=el.attrib.get('Sport','other').lower()
            if key=='Track' and points and points[-1] is not None:points.append(None)
            if key=='Id' and el.text and not times:times.append(el.text)
            if key=='Lap':
                for child in el:
                    if local_name(child.tag)=='TotalTimeSeconds':duration=(duration or 0)+float(child.text)
                    if local_name(child.tag)=='DistanceMeters':distance=(distance or 0)+float(child.text)
            if key=='Trackpoint':
                pos={}
                for child in el.iter():
                    k=local_name(child.tag)
                    if k in ('LatitudeDegrees','LongitudeDegrees'):pos[k]=float(child.text)
                    if k=='Time' and child.text:times.append(child.text)
                lat,lon=pos.get('LatitudeDegrees'),pos.get('LongitudeDegrees')
                if valid_point(lat,lon):points.append((lat,lon))
    start=None
    if times:
        dates=[datetime.fromisoformat(t.strip().replace('Z','+00:00')) for t in times]
        if any(d.tzinfo for d in dates) and not all(d.tzinfo for d in dates):
            # GPX/TCX times are UTC; do not fail on files mixing "…Z" and naive times.
            dates=[d if d.tzinfo else d.replace(tzinfo=timezone.utc) for d in dates]
        dates.sort()
        start=dates[0].isoformat()
        if duration is None:duration=max(0,(dates[-1]-dates[0]).total_seconds())
    sport={'biking':'cycling','bike':'cycling','ride':'cycling','run':'running','walk':'walking','hike':'hiking'}.get(sport,sport)
    summary={'activityName':name,'startTimeLocal':start,'activityType':{'typeKey':sport},
             'distance':distance if distance is not None else distance_m(points),'duration':duration}
    return summary,points

def parse_fit(data):
    import fitdecode
    points=[]; sessions=[]; timestamps=[]
    with fitdecode.FitReader(io.BytesIO(data),check_crc=fitdecode.CrcCheck.RAISE,error_handling=fitdecode.ErrorHandling.RAISE) as reader:
        for frame in reader:
            if not isinstance(frame,fitdecode.FitDataMessage):continue
            values={f.name:f.value for f in frame.fields}
            if frame.name=='record':
                lat,lon=values.get('position_lat'),values.get('position_long')
                if isinstance(lat,(int,float)) and isinstance(lon,(int,float)):
                    point=(lat*180/2**31,lon*180/2**31)
                    if valid_point(*point):points.append(point)
                if isinstance(values.get('timestamp'),datetime):timestamps.append(values['timestamp'])
            if frame.name=='session':sessions.append(values)
    if not sessions:raise ValueError('FIT file has no activity session (wellness files are not activity imports)')
    s=sessions[0];start=s.get('start_time') or (min(timestamps) if timestamps else None)
    sport=str(s.get('sport') or 'other')
    summary={'startTimeLocal':start.isoformat() if isinstance(start,datetime) else None,
             'activityType':{'typeKey':sport},'distance':s.get('total_distance'),'duration':s.get('total_timer_time'),
             'averageSpeed':s.get('enhanced_avg_speed') or s.get('avg_speed'),'averageHR':s.get('avg_heart_rate'),
             'elevationGain':s.get('total_ascent'),'elevationLoss':s.get('total_descent'),
             'calories':s.get('total_calories'),'avgPower':s.get('avg_power')}
    if len(sessions)>1:
        summary['activityType']={'typeKey':'multisport'}
        for dest,source in [('distance','total_distance'),('duration','total_timer_time'),('elevationGain','total_ascent'),('elevationLoss','total_descent'),('calories','total_calories')]:
            vals=[s[source] for s in sessions if isinstance(s.get(source),(int,float))]
            summary[dest]=sum(vals) if vals else None
        summary['averageHR']=None;summary['avgPower']=None
        summary['averageSpeed']=summary['distance']/summary['duration'] if summary.get('duration') and summary.get('distance') is not None else None
    return summary,points

def parse_activity(data,fmt):
    if len(data)>MAX_FILE:raise ValueError('File size limit exceeded')
    summary,points=parse_fit(data) if fmt=='fit' else parse_xml_activity(data,fmt)
    if not summary.get('startTimeLocal'):raise ValueError('Missing activity date')
    # Original files are kept whole. FIT multisport sessions are summarized together.
    return summary,points

def content_identity(a):
    # Same start, sport and near-identical distance/duration identify cross-format duplicates.
    stamp=a.get('startTimeGMT') or a.get('startTimeLocal') or ''
    try:
        dt=datetime.fromisoformat(stamp.replace('Z','+00:00'))
        if a.get('startTimeGMT') and not dt.tzinfo:dt=dt.replace(tzinfo=timezone.utc)
        stamp=dt.astimezone(timezone.utc).isoformat(timespec='seconds') if dt.tzinfo else dt.isoformat(timespec='seconds')
    except (ValueError,TypeError):pass
    values=[stamp,(a.get('activityType') or {}).get('typeKey'),round(a.get('distance') or 0),round(a.get('duration') or 0)]
    return hashlib.sha256(json.dumps(values,sort_keys=True).encode()).hexdigest()

def activities_from_archive(db, cancel=None):
    result={}
    rows=db.db.execute("SELECT * FROM responses WHERE dataset='activities' AND is_binary=0 ORDER BY fetched_utc")
    for row in rows:
        if cancel and cancel.is_set():raise Cancelled()
        entry=db.decode(row)
        for a in entry['data'] if isinstance(entry['data'],list) else []:
            if not isinstance(a,dict):continue
            a=dict(a);a['_day']=(a.get('startTimeLocal') or entry['day'])[:10]
            ident=str(a.get('activityId') or content_identity(a));a['_key']=ident
            result[ident]=a
    return list(result.values())

def track_from_archive(db,a):
    ident=a['_key'];prefix='activity-'+ident+'-'
    for suffix in ('gpx','tcx','fit','original'):
        rows=db.db.execute('SELECT * FROM responses WHERE dataset=? AND is_binary=1 ORDER BY fetched_utc DESC',(prefix+suffix,))
        for row in rows:
            data=bytes(row['payload'])
            try:
                if data[:2]==b'PK':
                    for name,fmt,part in activity_members(data):
                        try:return parse_activity(part,fmt)[1]
                        except (ValueError,ET.ParseError):continue
                else:return parse_fit(data)[1] if suffix=='fit' else parse_xml_activity(data,suffix)[1]
            except Exception:continue
    # Garmin details can contain GPS samples even without GPX.
    row=db.db.execute('SELECT * FROM responses WHERE dataset=? AND is_binary=0',(prefix+'details',)).fetchone()
    if row:
        payload=db.decode(row)['data']; descriptors=payload.get('metricDescriptors',[]) if isinstance(payload,dict) else []
        indexes={d.get('key'):d.get('metricsIndex') for d in descriptors}
        lat_i=indexes.get('directLatitude');lon_i=indexes.get('directLongitude');pts=[]
        if lat_i is not None and lon_i is not None:
            for p in payload.get('activityDetailMetrics',[]):
                v=p.get('metrics',[])
                if len(v)>max(lat_i,lon_i) and valid_point(v[lat_i],v[lon_i]):pts.append((v[lat_i],v[lon_i]))
        return pts
    return []

HEALTH_COLUMNS=[('date','Datum','Date'),('steps','Schritte','Steps'),('sleep_hours','Schlaf (h)','Sleep (h)'),
 ('sleep_score','Schlafscore','Sleep score'),('resting_hr','Ruhepuls (bpm)','Resting HR (bpm)'),
 ('hrv','HRV (ms)','HRV (ms)'),('stress','Ø Stress','Avg stress'),('body_battery_max','Body Battery Maximum','Body Battery maximum'),
 ('weight_kg','Gewicht (kg)','Weight (kg)'),('calories','Kalorien (kcal)','Calories (kcal)')]

def nested(value,path):
    for key in path.split('.'):
        if not isinstance(value,dict):return None
        value=value.get(key)
    return value

def daily_health(db,cancel=None):
    rows={}
    mappings={
      'summary':{'steps':('totalSteps',1),'resting_hr':('restingHeartRate',1),'stress':('averageStressLevel',1),
                 'body_battery_max':('bodyBatteryHighestValue',1),'calories':('totalKilocalories',1)},
      'heart_rate':{'resting_hr':('restingHeartRate',1)},
      'sleep':{'sleep_hours':('dailySleepDTO.sleepTimeSeconds',1/3600),'sleep_score':('dailySleepDTO.sleepScores.overall.value',1)},
      'hrv':{'hrv':('hrvSummary.lastNightAvg',1)},
      'stress':{'stress':('avgStressLevel',1)},
      'weight':{'weight_kg':('totalAverage.weight',.001)},
      'heart_rate-daily':{'resting_hr':('restingHeartRate',1)},
      'resting_hr-daily':{'resting_hr':('value',1)},
      'steps-daily':{'steps':('totalSteps',1)},
      'sleep-daily':{'sleep_hours':('totalSleepSeconds',1/3600),'sleep_score':('sleepScore',1)},
    }
    sql="SELECT * FROM responses WHERE dataset IN (%s) AND is_binary=0 AND day!='snapshot' ORDER BY fetched_utc"%(','.join('?' for _ in mappings))
    for raw in db.db.execute(sql,tuple(mappings)):
        if cancel and cancel.is_set():raise Cancelled()
        rec=db.decode(raw);day=rec['day'];dest=rows.setdefault(day,{'date':day})
        payload=rec['data']
        if rec['dataset'].endswith('-daily'):
            payload=payload[0] if isinstance(payload,list) and len(payload)==1 else {}
        for field,(path,scale) in mappings[rec['dataset']].items():
            value=nested(payload,path)
            if isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value) and value>=0:dest[field]=round(value*scale,2)
    return sorted(rows.values(),key=lambda r:r['date'])

def load_tracks(path,activities,cancel=None):
    tracks={};budget=max(200,2000000//max(1,len(activities)))
    with Archive(path,readonly=True) as db:
        for a in activities:
            if cancel and cancel.is_set():raise Cancelled()
            points=track_from_archive(db,a)
            if len(points)>budget:
                step=max(1,math.ceil(len(points)/budget))
                points=[p for i,p in enumerate(points) if p is None or i%step==0 or i==len(points)-1]
            tracks[a['_key']]=points
    return tracks

def load_archive(path,cancel=None,with_tracks=True):
    with Archive(path,readonly=True) as db:
        activities=activities_from_archive(db,cancel)
        tracks={}
        for a in activities if with_tracks else []:
            if cancel and cancel.is_set():raise Cancelled()
            tracks[a['_key']]=track_from_archive(db,a)
        import re
        candidates=set()
        for row in db.db.execute("SELECT DISTINCT dataset FROM responses WHERE dataset LIKE 'activity-%'"):
            match=re.fullmatch(r'activity-(-?\d+)-(gpx|tcx|fit|original|details)',row[0])
            if match:candidates.add(match.group(1))
        row=db.db.execute("SELECT COUNT(*),COALESCE(SUM(is_empty),0) FROM responses WHERE dataset='activities' AND is_binary=0").fetchone()
        failed=db.db.execute("SELECT COUNT(*) FROM attempts WHERE dataset='activities'").fetchone()[0]
        status={'responses':row[0],'empty_responses':row[1],'failed_requests':failed}
        return {'activities':activities,'tracks':tracks,'health':daily_health(db,cancel),'account':db.account,'track_candidates':candidates,'activity_status':status}

def activity_members(data,depth=0,budget=None):
    """Bounded in-memory traversal; never extract untrusted paths to disk."""
    if depth>4:raise ValueError('Archive nesting limit')
    budget=budget if budget is not None else [0,0]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for item in z.infolist():
            if item.is_dir():continue
            budget[0]+=item.file_size;budget[1]+=1
            if budget[0]>MAX_ZIP_TOTAL or budget[1]>MAX_ZIP_FILES:raise ValueError('Archive size limit')
            fmt=Path(item.filename).suffix.lower().lstrip('.')
            if fmt not in ('fit','gpx','tcx','zip'):continue
            if item.file_size>MAX_FILE:raise ValueError('Member size limit')
            part=z.read(item)
            if fmt=='zip':yield from activity_members(part,depth+1,budget)
            else:yield item.filename,fmt,part

def import_files(path,files,account,cancel=None,event=lambda *a:None):
    stats={'imported':0,'duplicates':0,'failed':0,'unsupported':0,'errors':[]}
    with Archive(path,account) as db:
        existing=activities_from_archive(db);known={content_identity(a):a for a in existing}
        by_day=defaultdict(list)
        for a in existing:
            if a.get('importSource'):by_day[a['_day']].append({k:v for k,v in a.items() if not k.startswith('_')})
        for filename in files:
            if cancel and cancel.is_set():break
            file=Path(filename)
            try:
                if file.stat().st_size>MAX_ZIP_TOTAL:raise ValueError('File size limit')
                data=file.read_bytes();fmt=file.suffix.lower().lstrip('.')
                members=activity_members(data) if fmt=='zip' else [(file.name,fmt,data)]
                for name,kind,part in members:
                    if cancel and cancel.is_set():break
                    try:
                        if kind not in ('fit','gpx','tcx'):stats['unsupported']+=1;continue
                        summary,points=parse_activity(part,kind);key=content_identity(summary)
                        if key in known:
                            ident=known[key]['activityId'];stats['duplicates']+=1
                        else:
                            ident=-int(hashlib.sha256(part).hexdigest()[:14],16)
                            summary.update(activityId=ident,activityName=summary.get('activityName') or Path(name).stem,
                                           importSource='local-'+kind)
                            if summary.get('duration') and summary.get('distance') is not None:
                                summary['averageSpeed']=summary['distance']/summary['duration']
                            day=summary['startTimeLocal'][:10];by_day[day].append(summary);known[key]=summary
                            # A separate import request avoids pretending a partial imported day is a full Garmin response.
                            db.put('activities',day,{'import':'activity-files-v1'},by_day[day])
                            stats['imported']+=1
                        day=known[key].get('_day') or known[key]['startTimeLocal'][:10]
                        db.put('activity-'+str(ident)+'-'+kind,day,{'import':'original-file-v1','format':kind},part,True)
                    except Exception as exc:
                        stats['failed']+=1
                        stats['errors'].append({'file':Path(name).name,'error_type':type(exc).__name__})
                    event('import',stats['imported'],stats['duplicates'],stats['failed'])
            except Exception as exc:
                stats['failed']+=1;stats['errors'].append({'file':file.name,'error_type':type(exc).__name__})
        stats['cancelled']=bool(cancel and cancel.is_set());db.record_run({'operation':'local_import',**stats})
    return stats

def aggregate(activities,period='month'):
    totals={}
    for a in activities:
        day=a.get('_day') or a.get('startTimeLocal','')[:10]
        key=(day[:4] if period=='year' else day[:7],(a.get('activityType') or {}).get('typeKey','other'))
        row=totals.setdefault(key,{'period':key[0],'sport':key[1],'count':0,'km':None,'hours':None,'ascent':None})
        row['count']+=1
        for name,source,scale in [('km','distance',.001),('hours','duration',1/3600),('ascent','elevationGain',1)]:
            val=a.get(source)
            if isinstance(val,(int,float)) and math.isfinite(val):row[name]=(row[name] or 0)+val*scale
    return [{k:round(v,2) if isinstance(v,float) else v for k,v in row.items()} for _,row in sorted(totals.items())]

def write_csv(path,rows,language='de'):
    fields=list(dict.fromkeys(k for r in rows for k in r))
    with open(path,'w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields,delimiter=';');writer.writeheader()
        for row in rows:writer.writerow({k:safe_cell(str(v).replace('.',',') if isinstance(v,float) and language=='de' else v) for k,v in row.items()})

def write_table(path,rows,language='de'):
    suffix=Path(path).suffix.lower()
    if suffix=='.json':
        Path(path).write_text(json.dumps(rows,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    elif suffix=='.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        from openpyxl.utils import get_column_letter
        wb=Workbook();sheet=wb.active;sheet.title='Data';fields=list(dict.fromkeys(k for r in rows for k in r))
        sheet.append(fields)
        for row in rows:sheet.append([safe_cell(row.get(k)) for k in fields])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087EAA')
        for i,key in enumerate(fields,1):sheet.column_dimensions[get_column_letter(i)].width=min(40,max(15,len(key)+3))
        wb.save(path)
    else:write_csv(path,rows,language)

def trim_track(points,meters):
    # Hide a radius around both endpoints; no connecting line across hidden areas.
    actual=[p for p in points if p is not None]
    if not actual or meters<=0:return points
    out=[]
    for p in points:
        if p is None or distance_m([actual[0],p])<meters or distance_m([actual[-1],p])<meters:
            if out and out[-1] is not None:out.append(None)
        else:out.append(p)
    return out

def write_kml(path,activities,tracks,privacy_m=0):
    ns='http://www.opengis.net/kml/2.2';ET.register_namespace('',ns)
    root=ET.Element('{'+ns+'}kml');doc=ET.SubElement(root,'Document')
    for a in activities:
        pts=trim_track(tracks.get(a['_key'],[]),privacy_m)
        groups=[];current=[]
        for p in pts+[None]:
            if p is None:
                if len(current)>1:groups.append(current)
                current=[]
            else:current.append(p)
        if not groups:continue
        mark=ET.SubElement(doc,'Placemark');ET.SubElement(mark,'name').text=a.get('activityName','Activity')
        geom=ET.SubElement(mark,'MultiGeometry')
        for group in groups:
            line=ET.SubElement(geom,'LineString');ET.SubElement(line,'tessellate').text='1'
            ET.SubElement(line,'coordinates').text=' '.join(f'{lon},{lat},0' for lat,lon in group)
    ET.ElementTree(root).write(path,encoding='utf-8',xml_declaration=True)

def backup_archive(source,destination):
    src=Path(source).resolve();dest=Path(destination).resolve()
    if src==dest:raise ValueError('Choose another destination')
    dest.parent.mkdir(parents=True,exist_ok=True)
    fd,name=tempfile.mkstemp(prefix='.garmin-backup-',suffix='.sqlite',dir=dest.parent);os.close(fd)
    try:
        # sqlite3's own context manager commits but does not close its handle.
        # Windows cannot rename the completed backup until the handle is closed.
        with Archive(src,readonly=True) as db, closing(sqlite3.connect(name)) as target:db.db.backup(target)
        validate_archive(name)
        os.replace(name,dest)
    finally:
        Path(name).unlink(missing_ok=True)
    return str(dest)

def validate_archive(path):
    with Archive(path,readonly=True) as db:
        result=db.db.execute('PRAGMA integrity_check').fetchone()[0]
        if result!='ok':raise OfflineArchiveError('Integrity check failed')
        return {'account':db.account,'responses':db.db.execute('SELECT COUNT(*) FROM responses').fetchone()[0],'integrity':'ok'}

def import_cache(path,cache_dir,account,cancel=None):
    """Import legacy exporter cache without re-downloading or treating CSV as raw data."""
    cache=Path(cache_dir).resolve()
    if cache.name!='cache-v1' or cache.parent.name!='account-'+account:raise OfflineArchiveError('Select the matching account cache-v1 folder')
    count=0;errors=0
    with Archive(path,account) as db:
        for file in cache.rglob('*'):
            if cancel and cancel.is_set():break
            if not file.is_file():continue
            rel=file.relative_to(cache)
            if len(rel.parts)!=3:continue
            dataset=rel.parts[0];day='snapshot' if file.name.startswith('snapshot-') else file.name[:10]
            try:
                if file.suffix=='.json':
                    envelope=json.loads(file.read_text(encoding='utf-8'))
                    if 'request' not in envelope or 'data' not in envelope:continue
                    db.put(dataset,day,envelope['request'],envelope['data'],False,envelope.get('fetched_utc'));count+=1
                elif dataset.startswith('activity-') and file.suffix in ('.gpx','.tcx','.zip','.kml','.csv'):
                    from garminconnect import Garmin
                    ident,fmt=dataset.removeprefix('activity-').rsplit('-',1)
                    if not ident.isdigit():continue
                    fmt=fmt.upper()
                    req=request_spec(dataset,'download_activity',(int(ident),),{'dl_fmt':Garmin.ActivityDownloadFormat[fmt]},fmt)
                    db.put(dataset,day,req,file.read_bytes(),True);count+=1
            except Exception:errors+=1
    return {'responses':count,'errors':errors}

def similar_routes(activities,tracks,cancel=None):
    """Approximate footprint matching on a 250m equatorial grid, no physiology inference."""
    from mapview import route_cells
    groups=[];grid=160000
    for a in activities:
        if cancel and cancel.is_set():raise Cancelled()
        cells=route_cells(tracks.get(a['_key'],[]),grid)
        if len(cells)<5:continue
        matched=False
        for group in groups:
            if len(cells)<len(group['cells'])*.8 or len(cells)>len(group['cells'])*1.25:continue
            if len(cells & group['cells'])/max(1,len(cells | group['cells']))>=.85:
                group['activities'].append(a);matched=True;break
        if not matched:groups.append({'cells':cells,'activities':[a]})
    return [{'trips':len(g['activities']),'example':g['activities'][0].get('activityName',''),
             'first':min(a['_day'] for a in g['activities']),'last':max(a['_day'] for a in g['activities'])}
            for g in groups if len(g['activities'])>1]
