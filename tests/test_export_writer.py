"""Conversion is streamed once; reuse must never hide changed/missing source data."""
from dataclasses import replace
from pathlib import Path
import csv
import json
import os
import pytest
from engine import Plan,account_key,request_spec
from offline import ArchiveExporter,OfflineExporter,ArchiveSync
from archive import Archive
from demo_data import DemoGarmin
from exportfiles import publish_export


def plan(tmp_path,**kw):
    return Plan(**(dict(start='2019-12-31',end='2020-01-01',datasets=['summary'],output=str(tmp_path/'out'),archive_path=str(tmp_path/'archive.sqlite'),
        account=account_key('writer@example.invalid'),formats=['csv','xlsx','json'],delay=0,refresh_days=0,optimize=False)|kw))


def test_three_formats_share_single_payload_pass(tmp_path,monkeypatch):
    p=plan(tmp_path);e=ArchiveExporter(DemoGarmin(),p);original=e.iter_payloads;counts=[]
    def counted(records):
        counts.extend(r['date'] for r in records)
        yield from original(records)
    monkeypatch.setattr(e,'iter_payloads',counted)
    dest,r=e.run();assert r['completed'] and r['output_stats']['written']==6
    assert counts==['2019-12-31','2020-01-01']
    assert r['timings']['write_seconds']>=0
    from openpyxl import load_workbook
    with (Path(dest)/'summary-2020.csv').open(encoding='utf-8-sig') as f:csv_rows=list(csv.reader(f,delimiter=';'))
    book=load_workbook(Path(dest)/'summary-2020.xlsx',read_only=True)
    xlsx_rows=list(book.active.values);book.close()
    assert csv_rows[0]==list(xlsx_rows[0]) and len(csv_rows)==len(xlsx_rows)
    assert json.loads((Path(dest)/'summary-2020.json').read_text(encoding='utf-8'))[0]['source_date']=='2020-01-01'


def test_repeat_reuses_files_without_row_conversion_or_replacement(tmp_path,monkeypatch):
    p=plan(tmp_path);dest,_=ArchiveExporter(DemoGarmin(),p).run();dest=Path(dest)
    before={f.name:f.stat().st_mtime_ns for f in dest.iterdir() if f.suffix in ('.csv','.xlsx')}
    e=OfflineExporter(replace(p,source='offline'))
    monkeypatch.setattr(e,'iter_payloads',lambda records:(_ for _ in ()).throw(AssertionError('reconverted')))
    result=e.run()[1]
    assert result['completed'] and result['output_stats']['reused']==6 and result['output_stats']['written']==0
    assert before=={f.name:f.stat().st_mtime_ns for f in dest.iterdir() if f.suffix in ('.csv','.xlsx')}


def test_only_changed_year_is_rewritten_and_force_rebuild_works(tmp_path):
    p=plan(tmp_path);dest,_=ArchiveExporter(DemoGarmin(),p).run()
    with Archive(p.archive_path,p.account) as db:
        db.put('summary','2020-01-01',request_spec('summary','get_user_summary',('2020-01-01',)),{'totalSteps':123456})
    report=OfflineExporter(replace(p,source='offline')).run()[1]
    assert report['output_stats']['reused']==3 and report['output_stats']['written']==3
    assert json.loads((Path(dest)/'summary-2020.json').read_text(encoding='utf-8'))[0]['data']['totalSteps']==123456
    report=OfflineExporter(replace(p,source='offline',reuse_exports=False)).run()[1]
    assert report['output_stats']['reused']==0 and report['output_stats']['written']==6


def test_modified_and_deleted_files_are_regenerated(tmp_path):
    p=plan(tmp_path,formats=['csv','json']);dest,_=ArchiveExporter(DemoGarmin(),p).run();dest=Path(dest)
    file=dest/'summary-2020.csv';before=file.read_bytes()
    file.write_bytes(b'X'+before[1:])  # same size still must be detected
    (dest/'summary-2019.json').unlink()
    report=OfflineExporter(replace(p,source='offline')).run()[1]
    assert report['output_stats']['written']==2 and report['output_stats']['reused']==2
    assert file.read_bytes()==before


def test_added_format_reuses_existing_and_changed_fields_invalidate(tmp_path):
    p=plan(tmp_path,formats=['json']);ArchiveExporter(DemoGarmin(),p).run()
    r=OfflineExporter(replace(p,source='offline',formats=['json','csv'])).run()[1]
    assert r['output_stats']['reused']==2 and r['output_stats']['written']==2
    r=OfflineExporter(replace(p,source='offline',formats=['csv'],selectors=['/totalSteps'])).run()[1]
    assert r['output_stats']['written']==2 and not r['output_stats']['reused']


def test_missing_source_never_uses_old_export_to_claim_completeness(tmp_path):
    p=plan(tmp_path,formats=['json']);ArchiveExporter(DemoGarmin(),p).run()
    with Archive(p.archive_path,p.account) as db:
        db.db.execute("DELETE FROM responses WHERE day='2020-01-01'");db.db.commit()
    dest,r=OfflineExporter(replace(p,source='offline')).run()
    assert not r['completed'] and r['stats']['errors']==1
    assert not (Path(dest)/'summary-2020.json').exists()


def test_wide_columns_json_and_tables_match(tmp_path):
    p=plan(tmp_path,layout='wide');dest,r=ArchiveExporter(DemoGarmin(),p).run()
    assert r['completed']
    with (Path(dest)/'summary-2020.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f,delimiter=';'))
    assert rows[0]['/totalSteps'] and '/restingHeartRate' in rows[0]


def test_retained_files_survive_publication_rollback(tmp_path,monkeypatch):
    dest=tmp_path/'current';dest.mkdir();(dest/'keep.csv').write_text('keep');(dest/'replace.csv').write_text('old')
    (dest/'report.json').write_text(json.dumps({'app':'ttn-garmin-export','generated_files':['keep.csv','replace.csv']}))
    stage=tmp_path/'stage';stage.mkdir();(stage/'replace.csv').write_text('new');(stage/'report.json').write_text('{}')
    original=os.replace
    def broken(a,b):
        if Path(a)==stage/'report.json':raise OSError('simulated failure')
        return original(a,b)
    monkeypatch.setattr(os,'replace',broken)
    with pytest.raises(OSError):publish_export(stage,dest,retained={'keep.csv'})
    assert (dest/'keep.csv').read_text()=='keep' and (dest/'replace.csv').read_text()=='old'


def test_download_only_never_creates_export_files(tmp_path,monkeypatch):
    p=plan(tmp_path);e=ArchiveSync(DemoGarmin(),p)
    monkeypatch.setattr(e,'write_outputs',lambda *a:(_ for _ in ()).throw(AssertionError('writer called')))
    r=e.run_sync()['report']
    assert r['completed'] and r['timings']['write_seconds']==0
    assert not list(Path(p.output).rglob('*.csv'))
    assert not list(Path(p.output).rglob('*.xlsx'))


def test_binary_files_are_retained_without_copying_again(tmp_path,monkeypatch):
    p=plan(tmp_path,start='2020-01-01',datasets=['activities'],formats=['json'],activity_files=['ORIGINAL'])
    class BinaryAPI(DemoGarmin):
        def download_activity(self,ident,dl_fmt):return b'original-file-'+str(ident).encode()
    dest,r=ArchiveExporter(BinaryAPI(),p).run();assert r['completed']
    binary=[x for x in r['records'] if x['binary']];assert binary
    e=OfflineExporter(replace(p,source='offline'))
    monkeypatch.setattr(e,'copy_binary',lambda *a:(_ for _ in ()).throw(AssertionError('unnecessary copy')))
    result=e.run()[1]
    assert result['completed'] and result['output_stats']['reused']==len(r['generated_files'])


def test_output_events_report_actual_steps_and_settings(tmp_path):
    events=[];p=plan(tmp_path)
    _,r=ArchiveExporter(DemoGarmin(),p,event=lambda *a:events.append(a)).run()
    assert any(e[0]=='settings' and e[1]['workers']==p.workers for e in events)
    assert [e[1] for e in events if e[0]=='phase_done']==['download','write']
    assert any(e[:3]==('output_step','start','xlsx_save') for e in events)
    assert any(e[0]=='output_progress' and e[1]['done']==6 and e[1]['total']==6 for e in events)
    assert r['output_stats']['steps']


def test_new_writer_matches_previous_csv_json_and_excel_values(tmp_path):
    # Independent reference expectations include locale, formula protection and nested values.
    from openpyxl import load_workbook
    class API:
        def get_user_summary(self,day):return {'value':1.25,'formula':'=1+1','nested':[None,False,{'label':'ä'}]}
    p=plan(tmp_path,start='2020-01-01');dest,r=ArchiveExporter(API(),p).run();assert r['completed']
    with (Path(dest)/'summary-2020.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f,delimiter=';'))
    by_path={x['path']:x for x in rows}
    assert by_path['/value']['value']=='1.25'
    assert by_path['/formula']['value']=="'=1+1"
    assert by_path['/nested/1']['value']=='false'
    book=load_workbook(Path(dest)/'summary-2020.xlsx',read_only=True)
    records=[dict(zip(next(iter(book.active.values)),v)) for v in list(book.active.values)[1:]];book.close()
    assert {x['path']:x['value'] for x in records}['/value']==1.25
    raw=json.loads((Path(dest)/'summary-2020.json').read_text(encoding='utf-8'))[0]['data']
    assert raw['formula']=='=1+1' and raw['nested']==[None,False,{'label':'ä'}]


def test_refreshed_identical_responses_keep_export_files(tmp_path):
    p=plan(tmp_path);ArchiveExporter(DemoGarmin(),p).run()
    _,r=ArchiveExporter(DemoGarmin(),replace(p,refresh_days=100000)).run()
    assert r['completed'] and r['stats']['downloaded']==2
    assert r['output_stats']['reused']==6 and r['output_stats']['written']==0
