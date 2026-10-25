"""Streaming output, content-based reuse, and measured local export progress."""
from contextlib import ExitStack
import csv
import hashlib
from branding import REPORT_APP_IDS
import json
from pathlib import Path
import re
import time
from exportfiles import safe_target

WRITER_VERSION=1


def file_hash(path):
    value=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):value.update(block)
    return value.hexdigest()


class OutputWriter:
    def __init__(self,exporter,destination):
        self.e=exporter;self.p=exporter.plan;self.destination=Path(destination)
        self.previous=getattr(exporter,'previous_output',self.destination)
        self.old={}
        if self.p.reuse_exports:
            try:
                report=json.loads(safe_target(self.previous,'report.json').read_text(encoding='utf-8'))
                if report.get('app') in REPORT_APP_IDS and report.get('config',{}).get('archive_path')==self.p.archive_path:
                    self.old=report.get('output_assets',{})
                    if not isinstance(self.old,dict):self.old={}
            except (OSError,ValueError,TypeError,AttributeError):pass
        self.completed=0;self.total=0
        exporter.retained_outputs=set();exporter.output_assets={}
        exporter.output_stats={'written':0,'reused':0,'steps':[]}

    def pulse(self,phase,name,rows=0,force=False):
        now=time.monotonic()
        if force or now-self.last_pulse>=.25:
            self.e.event('output_progress',{'phase':phase,'name':name,'rows':rows,'done':self.completed,'total':self.total,'seconds':now-self.started})
            self.last_pulse=now

    def begin(self,phase,name):
        self.started=time.monotonic();self.last_pulse=0
        self.e.event('output_step','start',phase,name,0.)
        self.pulse(phase,name,force=True)

    def end(self,phase,name):
        seconds=time.monotonic()-self.started
        self.e.output_stats['steps'].append({'phase':phase,'name':name,'seconds':seconds})
        self.e.event('output_step','end',phase,name,seconds)

    def signature(self,records,fmt):
        if any(not r.get('source_hash') for r in records):return None
        config={k:getattr(self.p,k) for k in ('account','layout','delimiter','split_year','activity_summary','language','selectors','activity_type')}
        # Sport hierarchy can change the effective activity filter without changing raw responses.
        config['taxonomy']={k:sorted(v) for k,v in (self.e.type_descendants or {}).items()} if self.p.activity_type else {}
        h=hashlib.sha256(json.dumps([WRITER_VERSION,fmt,config],sort_keys=True).encode())
        for r in records:h.update(json.dumps([r['dataset'],r['date'],r['source_hash']],ensure_ascii=False).encode())
        return h.hexdigest()

    def reuse(self,name,signature):
        if not signature:return False
        old=self.old.get(name)
        if not isinstance(old,dict) or old.get('signature')!=signature:return False
        try:
            path=safe_target(self.previous,name)
            if path.stat().st_size!=old.get('size'):return False
            self.begin('verify',name)
            matched=file_hash(path)==old.get('sha256')
            self.end('verify',name)
            if not matched:return False
        except (OSError,ValueError):return False
        self.e.retained_outputs.add(name);self.e.output_assets[name]=old
        self.e.output_stats['reused']+=1;self.completed+=1
        self.e.event('output_reused',name)
        return True

    def finish_file(self,name,signature):
        path=self.destination/name
        self.e.output_assets[name]={'signature':signature,'size':path.stat().st_size,'sha256':file_hash(path)}
        self.e.output_stats['written']+=1;self.completed+=1

    def run(self):
        groups={};binary=[]
        for record in self.e.records:
            if record['binary']:binary.append(record)
            else:
                year=record['date'][:4] if self.p.split_year else 'all'
                groups.setdefault(record['dataset']+'-'+year,[]).append(record)
        self.total=len(binary)+len(groups)*len(self.p.formats)
        for record in binary:
            folder='activities' if record['dataset'].startswith('activity-') else record['dataset']
            name=folder+'/'+Path(record['path']).name;signature=self.signature([record],'binary')
            if self.reuse(name,signature):continue
            self.begin('copy',name)
            target=self.destination/name;target.parent.mkdir(parents=True,exist_ok=True)
            self.e.copy_binary(record,target);self.finish_file(name,signature);self.end('copy',name)
        for key,records in groups.items():self.group(key,records)
        self.e.event('output_progress',{'phase':'complete','name':'','rows':0,'done':self.completed,'total':self.total,'seconds':0})

    def group(self,key,records):
        pending={fmt:self.signature(records,fmt) for fmt in self.p.formats}
        for fmt in list(pending):
            if self.reuse(key+'.'+fmt,pending[fmt]):del pending[fmt]
        if not pending:return
        table=bool(set(pending)&{'csv','xlsx'})
        summary=records[0]['dataset']=='activities' and self.p.activity_summary
        fields=self.e.summary_fields() if summary else (['source_date','dataset','path','value','type'] if self.p.layout=='long' else ['source_date','dataset','record_index'])
        if table and self.p.layout=='wide' and not summary:
            self.begin('columns',key);found=set(fields)
            for i,row in enumerate(self.e.rows(records)):
                for name in row:
                    if name not in found:fields.append(name);found.add(name)
                if i%1000==0:self.pulse('columns',key,i)
            self.end('columns',key)
        if 'xlsx' in pending and len(fields)>16384:raise ValueError('Excel column limit; use CSV/JSON')
        phase='+'.join(sorted(pending));self.begin(phase,key)
        from engine import safe_cell
        book=None;sheet=None;n=1;page=1;row_count=0
        try:
            with ExitStack() as stack:
                jf=stack.enter_context((self.destination/(key+'.json')).open('w',encoding='utf-8')) if 'json' in pending else None
                if jf:jf.write('[\n')
                writer=None
                if 'csv' in pending:
                    cf=stack.enter_context((self.destination/(key+'.csv')).open('w',encoding='utf-8-sig',newline=''))
                    writer=csv.DictWriter(cf,fieldnames=fields,delimiter=self.p.delimiter);writer.writeheader()
                if 'xlsx' in pending:
                    from openpyxl import Workbook
                    book=Workbook(write_only=True);sheet=book.create_sheet('Data 1');sheet.append(fields)
                for i,(record,payload) in enumerate(self.e.iter_payloads(records)):
                    if jf:
                        if i:jf.write(',\n')
                        # Encode one response at a time with the C encoder; never buffer the full history.
                        jf.write(json.dumps({'source_date':record['date'],'dataset':record['dataset'],'data':payload},ensure_ascii=False))
                    if table:
                        for row in self.e.rows_from_payload(record,payload):
                            if writer:writer.writerow(self.e.table_cells(row,summary))
                            if book:
                                if n>=1048576:
                                    page+=1;sheet=book.create_sheet('Data '+str(page));sheet.append(fields);n=1
                                values=[safe_cell(row.get(k)) for k in fields]
                                if any(isinstance(v,str) and (len(v)>32767 or re.search(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]',v)) for v in values):
                                    raise ValueError('Value cannot be represented in Excel; use CSV/JSON')
                                sheet.append(values);n+=1
                            row_count+=1
                            if row_count%1000==0:self.pulse(phase,key,row_count)
                    self.pulse(phase,key,row_count)
                if jf:jf.write('\n]\n')
            self.end(phase,key)
            if book:
                self.begin('xlsx_save',key);book.save(self.destination/(key+'.xlsx'));self.end('xlsx_save',key)
        finally:
            if book:
                for s in book.worksheets:
                    if not s.closed:s.close()
                book.close()
        self.begin('checksum',key)
        for fmt,signature in pending.items():self.finish_file(key+'.'+fmt,signature)
        self.end('checksum',key)


def write_outputs(exporter,destination):
    OutputWriter(exporter,destination).run()
